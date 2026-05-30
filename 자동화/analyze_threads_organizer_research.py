from __future__ import annotations

from pathlib import Path
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


RUN_STEM = "threads_organizer_research_20260514_173459"


COL = {
    "keyword": "검색어",
    "author": "작성자",
    "url": "게시글링크",
    "body": "본문",
    "body_block": "본문_후보블록",
    "likes": "좋아요수",
    "comments": "댓글수",
    "reposts": "리포스트수",
    "shares": "공유수",
    "engagement": "반응합계",
    "has_link": "외부링크있음",
    "link": "외부링크",
    "media_count": "네이티브미디어수",
    "recruit_score": "모집점수",
    "recruit_hits": "모집히트",
    "meetup_hits": "모임히트",
    "noise_hits": "잡음히트",
    "label_status": "라벨상태",
    "author_type": "작성자유형",
    "post_type": "글유형",
    "category": "모임카테고리",
    "performance_label": "성과라벨",
    "include_training": "학습포함",
    "exclude_reason": "제외사유",
    "label_note": "라벨메모",
}


def summarize(value: object, limit: int = 120) -> str:
    text = " ".join(str(value or "").split())
    if len(text) <= limit:
        return text
    return text[: limit - 1] + "…"


def comma_join(values: pd.Series) -> str:
    clean = [str(v).strip() for v in values if str(v).strip() and str(v) != "nan"]
    return ", ".join(sorted(set(clean)))


def joined_samples(values: pd.Series, limit: int = 3) -> str:
    clean = [summarize(v, 90) for v in values if str(v).strip() and str(v) != "nan"]
    return " / ".join(clean[:limit])


def add_model_features(df: pd.DataFrame, source_name: str) -> pd.DataFrame:
    out = df.copy()
    for col in [COL["performance_label"], COL["exclude_reason"], COL["label_note"]]:
        if col in out.columns:
            out[col] = out[col].astype("object")
    if COL["include_training"] in out.columns:
        out[COL["include_training"]] = out[COL["include_training"]].astype("object")

    body_summary = out[COL["body"]].map(lambda v: summarize(v, 180))
    out["본문요약"] = body_summary

    recruit_score = pd.to_numeric(out[COL["recruit_score"]], errors="coerce").fillna(0)
    engagement = pd.to_numeric(out[COL["engagement"]], errors="coerce").fillna(0)
    has_recruit = out[COL["recruit_hits"]].fillna("").astype(str).str.len().gt(0)
    has_meetup = out[COL["meetup_hits"]].fillna("").astype(str).str.len().gt(0)
    has_noise = out[COL["noise_hits"]].fillna("").astype(str).str.len().gt(0)

    out["작성자후보적합"] = (recruit_score >= 6) & has_recruit & has_meetup & ~has_noise
    out["모집글적합"] = (recruit_score >= 4) & has_recruit & has_meetup & ~has_noise
    out["성과점수_raw"] = engagement
    out["성과분위_source"] = source_name
    out["성과상위20"] = False
    out["성과하위40"] = False

    eligible = out[out["모집글적합"]].copy()
    if not eligible.empty:
        q80 = eligible["성과점수_raw"].quantile(0.80)
        q40 = eligible["성과점수_raw"].quantile(0.40)
        out.loc[out["모집글적합"] & (out["성과점수_raw"] >= q80), "성과상위20"] = True
        out.loc[out["모집글적합"] & (out["성과점수_raw"] <= q40), "성과하위40"] = True
        out.loc[out["성과상위20"], COL["performance_label"]] = "top_discovery"
        out.loc[out["성과하위40"], COL["performance_label"]] = "low_discovery"
        out.loc[out["모집글적합"], COL["include_training"]] = True
        out.loc[out["모집글적합"], COL["exclude_reason"]] = ""

    return out


def build_strict_handles(discovery: pd.DataFrame) -> pd.DataFrame:
    strict = discovery[discovery["작성자후보적합"]].copy()
    if strict.empty:
        return pd.DataFrame(
            columns=[
                COL["author"],
                "강한후보글수",
                "평균모집점수",
                "최대반응합계",
                "검색어목록",
                "대표근거",
            ]
        )

    grouped = (
        strict.groupby(COL["author"], dropna=False)
        .agg(
            강한후보글수=(COL["url"], "count"),
            평균모집점수=(COL["recruit_score"], "mean"),
            최대반응합계=(COL["engagement"], "max"),
            검색어목록=(COL["keyword"], comma_join),
            대표근거=("본문요약", joined_samples),
        )
        .reset_index()
    )
    grouped["평균모집점수"] = grouped["평균모집점수"].round(1)
    return grouped.sort_values(["강한후보글수", "평균모집점수", "최대반응합계"], ascending=False)


def build_profile_summary(profiles: pd.DataFrame) -> pd.DataFrame:
    if profiles.empty:
        return pd.DataFrame()
    grouped = (
        profiles.groupby(COL["author"], dropna=False)
        .agg(
            프로필글수=(COL["url"], "count"),
            평균반응합계=(COL["engagement"], "mean"),
            최대반응합계=(COL["engagement"], "max"),
            top글수=("계정내성과라벨", lambda s: int((s == "top").sum())),
            low글수=("계정내성과라벨", lambda s: int((s == "low").sum())),
            모집글수=(COL["recruit_score"], lambda s: int((pd.to_numeric(s, errors="coerce").fillna(0) >= 3).sum())),
        )
        .reset_index()
    )
    grouped["평균반응합계"] = grouped["평균반응합계"].round(1)
    return grouped.sort_values(["모집글수", "최대반응합계"], ascending=False)


def build_notes(quality: pd.DataFrame, discovery: pd.DataFrame, profiles: pd.DataFrame, strict_handles: pd.DataFrame) -> pd.DataFrame:
    recruit_ready = int(discovery["모집글적합"].sum())
    author_ready = int(discovery["작성자후보적합"].sum())
    profile_counts = profiles.groupby(COL["author"])[COL["url"]].count() if not profiles.empty else pd.Series(dtype=int)
    max_profile_posts = int(profile_counts.max()) if not profile_counts.empty else 0

    rows = [
        {
            "항목": "판단",
            "값": (
                "프로필 기반 수집은 비로그인 상태에서 작성자당 공개 카드가 4개 안팎이라 "
                "계정별 ML 학습 데이터로는 부족합니다. 이번 산출물은 discovery 검색글을 "
                "모집/모임 적합 글로 라벨링해 ML 원천 데이터로 쓰는 쪽이 더 맞습니다."
            ),
        },
        {"항목": "ML 우선 시트", "값": "discovery_model_ready"},
        {"항목": "운영자 후보 시트", "값": "strict_candidate_handles"},
        {"항목": "수집 규모", "값": f"검색글 {len(discovery)}개, 프로필글 {len(profiles)}개, 강한 작성자 후보 {len(strict_handles)}명"},
        {"항목": "학습 후보 규모", "값": f"모집글 적합 {recruit_ready}개, 작성자후보 적합 {author_ready}개"},
        {"항목": "프로필 한계", "값": f"프로필별 최대 공개 수집 글수 {max_profile_posts}개"},
    ]
    for _, row in quality.iterrows():
        rows.append({"항목": f"원본_{row.get('항목')}", "값": row.get("값")})
    return pd.DataFrame(rows)


def format_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    widths = {
        "quality_summary": {"A": 24, "B": 110},
        "research_notes": {"A": 24, "B": 120},
        "strict_candidate_handles": {"A": 24, "E": 44, "F": 110},
        "profile_summary": {"A": 24},
        "candidate_handles": {"A": 24, "E": 110, "F": 80},
        "profile_posts_ml": {"H": 90, "I": 80, "AO": 80},
        "discovery_posts": {"H": 90, "I": 80, "AD": 80},
        "discovery_model_ready": {"H": 90, "I": 80, "AD": 80, "AE": 80},
    }
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row and ws.max_column:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        for col, width in widths.get(ws.title, {}).items():
            ws.column_dimensions[col].width = width
    wb.save(path)


def main() -> None:
    run_dir = Path(__file__).resolve().parent / "organizer_research_runs"
    path = run_dir / f"{RUN_STEM}.xlsx"
    sheets = pd.read_excel(path, sheet_name=None)

    quality = sheets["quality_summary"]
    handles = sheets["candidate_handles"].copy()
    profiles = add_model_features(sheets["profile_posts_ml"], "profile")
    discovery = add_model_features(sheets["discovery_posts"], "discovery")

    handles["근거요약"] = handles["근거본문"].map(lambda v: summarize(v, 180))
    strict_handles = build_strict_handles(discovery)
    profile_summary = build_profile_summary(profiles)
    notes = build_notes(quality, discovery, profiles, strict_handles)

    priority_cols = [
        COL["keyword"],
        COL["author"],
        COL["url"],
        COL["body"],
        "본문요약",
        COL["likes"],
        COL["comments"],
        COL["reposts"],
        COL["shares"],
        COL["engagement"],
        COL["has_link"],
        COL["link"],
        COL["media_count"],
        COL["recruit_score"],
        COL["recruit_hits"],
        COL["meetup_hits"],
        COL["noise_hits"],
        "모집글적합",
        "작성자후보적합",
        "성과상위20",
        "성과하위40",
        COL["performance_label"],
        COL["include_training"],
        COL["exclude_reason"],
        COL["category"],
        COL["post_type"],
        COL["author_type"],
        COL["label_note"],
    ]
    model_ready = discovery[[c for c in priority_cols if c in discovery.columns]].copy()
    model_ready = model_ready.sort_values(["모집글적합", COL["recruit_score"], COL["engagement"]], ascending=False)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        quality.to_excel(writer, sheet_name="quality_summary", index=False)
        notes.to_excel(writer, sheet_name="research_notes", index=False)
        strict_handles.to_excel(writer, sheet_name="strict_candidate_handles", index=False)
        profile_summary.to_excel(writer, sheet_name="profile_summary", index=False)
        model_ready.to_excel(writer, sheet_name="discovery_model_ready", index=False)
        handles.to_excel(writer, sheet_name="candidate_handles", index=False)
        profiles.to_excel(writer, sheet_name="profile_posts_ml", index=False)
        discovery.to_excel(writer, sheet_name="discovery_posts", index=False)
        sheets["seed_summary"].to_excel(writer, sheet_name="seed_summary", index=False)
        sheets["run_config"].to_excel(writer, sheet_name="run_config", index=False)

    format_workbook(path)

    print(f"updated={path}")
    print("\nresearch_notes")
    print(notes.to_string(index=False))
    print("\nstrict_candidate_handles")
    print(strict_handles.head(15).to_string(index=False))
    print("\nprofile_summary")
    print(profile_summary.head(15).to_string(index=False))
    print("\ndiscovery_model_ready_top")
    print(
        model_ready.head(20)[
            [
                COL["keyword"],
                COL["author"],
                COL["recruit_score"],
                COL["engagement"],
                "모집글적합",
                "작성자후보적합",
                "본문요약",
            ]
        ].to_string(index=False)
    )


if __name__ == "__main__":
    main()
