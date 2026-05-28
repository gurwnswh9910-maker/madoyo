from __future__ import annotations

from pathlib import Path
import json
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


RUN_DIR = Path(__file__).resolve().parent / "organizer_research_runs"
INPUT_XLSX = RUN_DIR / "threads_meetup_profiles_relationship_labeled_20260515.xlsx"
OUTPUT_XLSX = RUN_DIR / "threads_manual_label_needed_20260515.xlsx"
OUTPUT_JSONL = RUN_DIR / "threads_manual_label_needed_20260515.jsonl"


def clean(value: object) -> str:
    text = " ".join(str(value or "").split())
    return "" if text.lower() == "nan" else text


def has_both_hits(row: pd.Series) -> bool:
    return bool(clean(row.get("학습히트")) and clean(row.get("관계히트")))


def is_medium_or_low(row: pd.Series) -> bool:
    return clean(row.get("라벨신뢰도")) in {"medium", "low"}


def is_review_label(row: pd.Series) -> bool:
    return clean(row.get("목적라벨")) == "검토" or clean(row.get("최종목적라벨")) == "검토"


def has_profile_evidence_conflict(row: pd.Series) -> bool:
    profile_label = clean(row.get("목적라벨"))
    final_label = clean(row.get("최종목적라벨"))
    if not profile_label or not final_label:
        return False
    return profile_label != final_label and final_label in {"관계", "학습"} and profile_label in {"관계", "학습", "검토"}


def close_evidence_count(row: pd.Series) -> bool:
    relation_count = int(row.get("근거글_관계수") or 0)
    learning_count = int(row.get("근거글_학습수") or 0)
    if relation_count == 0 and learning_count == 0:
        return False
    return abs(relation_count - learning_count) <= 1 and min(relation_count, learning_count) > 0


def ambiguous_reason(row: pd.Series) -> str:
    reasons = []
    if is_review_label(row):
        reasons.append("검토 라벨")
    if is_medium_or_low(row):
        reasons.append(f"신뢰도 {clean(row.get('라벨신뢰도'))}")
    if has_both_hits(row):
        reasons.append("관계/학습 히트 동시 존재")
    if has_profile_evidence_conflict(row):
        reasons.append(f"프로필라벨({clean(row.get('목적라벨'))}) vs 근거기반({clean(row.get('최종목적라벨'))}) 충돌")
    if close_evidence_count(row):
        reasons.append(f"근거글 관계/학습 근접({int(row.get('근거글_관계수') or 0)}:{int(row.get('근거글_학습수') or 0)})")
    return " / ".join(reasons)


def add_manual_columns(df: pd.DataFrame, source_sheet: str) -> pd.DataFrame:
    out = df.copy()
    out.insert(0, "수동라벨", "")
    out.insert(1, "수동판단메모", "")
    out.insert(2, "최종사용여부", "")
    out.insert(3, "사람판단필요사유", out.apply(ambiguous_reason, axis=1))
    out.insert(4, "source_sheet", source_sheet)
    return out


def pick_profile_cases(profiles: pd.DataFrame) -> pd.DataFrame:
    mask = profiles.apply(
        lambda row: (
            is_review_label(row)
            or is_medium_or_low(row)
            or has_both_hits(row)
            or has_profile_evidence_conflict(row)
            or close_evidence_count(row)
        ),
        axis=1,
    )
    picked = profiles[mask].copy()
    return add_manual_columns(picked, "profiles_labeled")


def pick_evidence_cases(evidence: pd.DataFrame) -> pd.DataFrame:
    mask = evidence.apply(
        lambda row: is_review_label(row) or is_medium_or_low(row) or has_both_hits(row),
        axis=1,
    )
    picked = evidence[mask].copy()
    return add_manual_columns(picked, "evidence_posts_labeled")


def pick_raw_cases(raw: pd.DataFrame) -> pd.DataFrame:
    # Raw is large, so keep only rows the rule engine explicitly marked as uncertain.
    mask = raw.apply(
        lambda row: (
            is_review_label(row)
            or is_medium_or_low(row)
        )
        and clean(row.get("모집글적합")) != "False",
        axis=1,
    )
    picked = raw[mask].copy()
    return add_manual_columns(picked, "raw_posts_labeled")


def format_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="7F6000")
    header_font = Font(color="FFFFFF", bold=True)
    manual_fill = PatternFill("solid", fgColor="FFF2CC")

    widths = {
        "quality_summary": {"A": 30, "B": 90},
        "manual_profile_review": {"A": 14, "B": 34, "C": 14, "D": 60, "E": 26, "J": 24, "K": 38, "W": 90, "X": 90},
        "manual_evidence_review": {"A": 14, "B": 34, "C": 14, "D": 60, "E": 26, "N": 36, "O": 90},
        "manual_raw_review": {"A": 14, "B": 34, "C": 14, "D": 60, "E": 26, "Q": 36, "R": 90},
    }

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row and ws.max_column:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=3):
            for cell in row:
                cell.fill = manual_fill
        for col, width in widths.get(ws.title, {}).items():
            ws.column_dimensions[col].width = width

        if ws.title.startswith("manual_") and ws.max_row >= 2:
            label_validation = DataValidation(type="list", formula1='"관계,학습,제외,보류"', allow_blank=True)
            use_validation = DataValidation(type="list", formula1='"Y,N,보류"', allow_blank=True)
            ws.add_data_validation(label_validation)
            ws.add_data_validation(use_validation)
            label_validation.add(f"A2:A{ws.max_row}")
            use_validation.add(f"C2:C{ws.max_row}")

    wb.save(path)


def main() -> None:
    sheets = pd.read_excel(INPUT_XLSX, sheet_name=None)
    profiles = sheets["profiles_labeled"].copy()
    evidence = sheets["evidence_posts_labeled"].copy()
    raw = sheets["raw_posts_labeled"].copy()

    profile_cases = pick_profile_cases(profiles)
    evidence_cases = pick_evidence_cases(evidence)
    raw_cases = pick_raw_cases(raw)

    quality = pd.DataFrame(
        [
            {"항목": "input", "값": str(INPUT_XLSX)},
            {"항목": "output_excel", "값": str(OUTPUT_XLSX)},
            {"항목": "output_jsonl", "값": str(OUTPUT_JSONL)},
            {"항목": "manual_profile_review_rows", "값": len(profile_cases)},
            {"항목": "manual_evidence_review_rows", "값": len(evidence_cases)},
            {"항목": "manual_raw_review_rows", "값": len(raw_cases)},
            {
                "항목": "선별 기준",
                "값": "검토/medium/low/관계-학습 히트 동시/프로필-근거 라벨 충돌/근거글 관계-학습 수 근접",
            },
            {
                "항목": "수동 입력 컬럼",
                "값": "수동라벨(관계, 학습, 제외, 보류), 수동판단메모, 최종사용여부(Y, N, 보류)",
            },
        ]
    )

    with OUTPUT_JSONL.open("w", encoding="utf-8") as handle:
        for row in profile_cases.to_dict(orient="records"):
            handle.write(json.dumps({"level": "profile", **row}, ensure_ascii=False, default=str) + "\n")
        for row in evidence_cases.to_dict(orient="records"):
            handle.write(json.dumps({"level": "evidence", **row}, ensure_ascii=False, default=str) + "\n")
        for row in raw_cases.to_dict(orient="records"):
            handle.write(json.dumps({"level": "raw", **row}, ensure_ascii=False, default=str) + "\n")

    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        quality.to_excel(writer, sheet_name="quality_summary", index=False)
        profile_cases.to_excel(writer, sheet_name="manual_profile_review", index=False)
        evidence_cases.to_excel(writer, sheet_name="manual_evidence_review", index=False)
        raw_cases.to_excel(writer, sheet_name="manual_raw_review", index=False)

    format_workbook(OUTPUT_XLSX)

    print(f"output_excel={OUTPUT_XLSX}")
    print(f"output_jsonl={OUTPUT_JSONL}")
    print("\nquality_summary")
    print(quality.to_string(index=False))
    print("\nmanual_profile_review")
    if profile_cases.empty:
        print("(empty)")
    else:
        cols = [c for c in ["q_id", "사람판단필요사유", "작성자", "profile_url", "목적라벨", "최종목적라벨", "라벨근거", "대표근거"] if c in profile_cases.columns]
        print(profile_cases[cols].head(40).to_string(index=False))


if __name__ == "__main__":
    main()
