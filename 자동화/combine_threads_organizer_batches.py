from __future__ import annotations

from pathlib import Path
import json
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


BATCH_STEMS = [
    "threads_organizer_research_20260514_181805",
    "threads_organizer_research_20260514_182119",
    "threads_organizer_research_20260514_182434",
    "threads_organizer_research_20260514_182855",
    "threads_organizer_research_20260514_183247",
    "threads_organizer_research_20260514_183545",
    "threads_organizer_research_20260514_184309",
]

OUTPUT_STEM = "threads_organizer_ml_discovery_combined_20260514"

COL = {
    "keyword": "검색어",
    "author": "작성자",
    "url": "게시글링크",
    "body": "본문",
    "likes": "좋아요수",
    "comments": "댓글수",
    "reposts": "리포스트수",
    "shares": "공유수",
    "engagement": "반응합계",
    "has_link": "외부링크있음",
    "link": "외부링크",
    "media_count": "네이티브미디어수",
    "recruit_score": "모집점수",
    "recruit_hits": "모집히트",
    "meetup_hits": "모임히트",
    "noise_hits": "잡음히트",
    "category": "모임카테고리",
    "performance_label": "성과라벨",
    "include_training": "학습포함",
    "exclude_reason": "제외사유",
}


def summarize(value: object, limit: int = 180) -> str:
    text = " ".join(str(value or "").split())
    if len(text) <= limit:
        return text
    return text[: limit - 1] + "…"


def clean_text(value: object) -> str:
    text = str(value or "").strip()
    return "" if text.lower() == "nan" else text


def comma_join(values: pd.Series) -> str:
    clean = [clean_text(v) for v in values if clean_text(v)]
    return ", ".join(sorted(set(clean)))


def sample_join(values: pd.Series, limit: int = 3) -> str:
    clean = [summarize(v, 100) for v in values if clean_text(v)]
    return " / ".join(clean[:limit])


def add_labels(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out["본문요약"] = out[COL["body"]].map(summarize)

    recruit_score = pd.to_numeric(out[COL["recruit_score"]], errors="coerce").fillna(0)
    engagement = pd.to_numeric(out[COL["engagement"]], errors="coerce").fillna(0)
    has_recruit = out[COL["recruit_hits"]].fillna("").astype(str).str.strip().ne("")
    has_meetup = out[COL["meetup_hits"]].fillna("").astype(str).str.strip().ne("")
    has_noise = out[COL["noise_hits"]].fillna("").astype(str).str.strip().ne("")

    out["모집글적합"] = (recruit_score >= 4) & has_recruit & has_meetup & ~has_noise
    out["작성자후보적합"] = (recruit_score >= 6) & has_recruit & has_meetup & ~has_noise
    out["성과점수_raw"] = engagement
    out["성과상위20_global"] = False
    out["성과하위40_global"] = False
    out["성과상위20_keyword"] = False
    out["성과하위40_keyword"] = False
    out["성과라벨_global"] = ""
    out["성과라벨_keyword"] = ""
    out["binary_label_global"] = ""
    out["binary_label_keyword"] = ""
    out["학습권장"] = False
    out["학습권장사유"] = ""

    eligible = out[out["모집글적합"]].copy()
    if not eligible.empty:
        q80 = eligible["성과점수_raw"].quantile(0.80)
        q40 = eligible["성과점수_raw"].quantile(0.40)
        top_global = out["모집글적합"] & (out["성과점수_raw"] >= q80)
        low_global = out["모집글적합"] & (out["성과점수_raw"] <= q40)
        out.loc[top_global, "성과상위20_global"] = True
        out.loc[low_global, "성과하위40_global"] = True
        out.loc[top_global, "성과라벨_global"] = "top"
        out.loc[low_global, "성과라벨_global"] = "low"
        out.loc[top_global, "binary_label_global"] = 1
        out.loc[low_global, "binary_label_global"] = 0

        for _, group in eligible.groupby(COL["keyword"]):
            if len(group) < 5:
                continue
            q80_kw = group["성과점수_raw"].quantile(0.80)
            q40_kw = group["성과점수_raw"].quantile(0.40)
            indexes = group.index
            kw_top = indexes[group["성과점수_raw"] >= q80_kw]
            kw_low = indexes[group["성과점수_raw"] <= q40_kw]
            out.loc[kw_top, "성과상위20_keyword"] = True
            out.loc[kw_low, "성과하위40_keyword"] = True
            out.loc[kw_top, "성과라벨_keyword"] = "top"
            out.loc[kw_low, "성과라벨_keyword"] = "low"
            out.loc[kw_top, "binary_label_keyword"] = 1
            out.loc[kw_low, "binary_label_keyword"] = 0

    recommended = out["모집글적합"] & (
        out["성과라벨_global"].isin(["top", "low"]) | out["성과라벨_keyword"].isin(["top", "low"])
    )
    out.loc[recommended, "학습권장"] = True
    out.loc[recommended, "학습권장사유"] = "모집/모임 적합 + global 또는 keyword top/low 라벨"
    out.loc[out["모집글적합"] & ~recommended, "학습권장사유"] = "중간 성과: 회귀/랭킹용 raw로 유지"
    out.loc[~out["모집글적합"], "학습권장사유"] = "모집/모임 적합 조건 미달 또는 잡음 히트"
    return out


def build_strict_handles(df: pd.DataFrame) -> pd.DataFrame:
    strict = df[df["작성자후보적합"]].copy()
    if strict.empty:
        return pd.DataFrame()
    grouped = (
        strict.groupby(COL["author"], dropna=False)
        .agg(
            강한후보글수=(COL["url"], "count"),
            평균모집점수=(COL["recruit_score"], "mean"),
            최대반응합계=(COL["engagement"], "max"),
            평균반응합계=(COL["engagement"], "mean"),
            검색어목록=(COL["keyword"], comma_join),
            대표근거=("본문요약", sample_join),
        )
        .reset_index()
    )
    grouped["평균모집점수"] = grouped["평균모집점수"].round(1)
    grouped["평균반응합계"] = grouped["평균반응합계"].round(1)
    return grouped.sort_values(["강한후보글수", "최대반응합계", "평균모집점수"], ascending=False)


def format_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    widths = {
        "research_notes": {"A": 24, "B": 120},
        "quality_summary": {"A": 28, "B": 80},
        "ml_training_candidates": {"A": 16, "B": 22, "C": 55, "D": 90, "E": 70},
        "top_posts": {"A": 16, "B": 22, "C": 55, "D": 90, "E": 70},
        "low_posts": {"A": 16, "B": 22, "C": 55, "D": 90, "E": 70},
        "strict_candidate_handles": {"A": 24, "F": 48, "G": 110},
        "seed_summary": {"A": 22},
        "discovery_posts_raw": {"H": 90, "I": 80},
        "source_files": {"A": 90},
    }
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row and ws.max_column:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        for col, width in widths.get(ws.title, {}).items():
            ws.column_dimensions[col].width = width
    wb.save(path)


def main() -> None:
    run_dir = Path(__file__).resolve().parent / "organizer_research_runs"
    frames = []
    seed_frames = []
    sources = []
    for stem in BATCH_STEMS:
        path = run_dir / f"{stem}.xlsx"
        sheets = pd.read_excel(path, sheet_name=None)
        discovery = sheets["discovery_posts"].copy()
        discovery["source_file"] = path.name
        frames.append(discovery)
        seed_summary = sheets["seed_summary"].copy()
        seed_summary["source_file"] = path.name
        seed_frames.append(seed_summary)
        sources.append({"source_file": str(path), "rows": len(discovery)})

    raw = pd.concat(frames, ignore_index=True)
    raw_before = len(raw)
    raw = raw.drop_duplicates(subset=[COL["url"]], keep="last").reset_index(drop=True)
    labeled = add_labels(raw)
    strict_handles = build_strict_handles(labeled)
    seed_summary = pd.concat(seed_frames, ignore_index=True)

    priority_cols = [
        COL["keyword"],
        COL["author"],
        COL["url"],
        COL["body"],
        "본문요약",
        COL["likes"],
        COL["comments"],
        COL["reposts"],
        COL["shares"],
        COL["engagement"],
        COL["has_link"],
        COL["link"],
        COL["media_count"],
        COL["recruit_score"],
        COL["recruit_hits"],
        COL["meetup_hits"],
        COL["noise_hits"],
        "모집글적합",
        "작성자후보적합",
        "성과라벨_global",
        "binary_label_global",
        "성과라벨_keyword",
        "binary_label_keyword",
        "학습권장",
        "학습권장사유",
        "source_file",
    ]
    candidates = labeled[[c for c in priority_cols if c in labeled.columns]].copy()
    candidates = candidates.sort_values(["학습권장", "모집글적합", COL["recruit_score"], COL["engagement"]], ascending=False)
    training = candidates[candidates["학습권장"]].copy()
    top_posts = training[(training["성과라벨_global"] == "top") | (training["성과라벨_keyword"] == "top")].copy()
    low_posts = training[(training["성과라벨_global"] == "low") | (training["성과라벨_keyword"] == "low")].copy()

    quality = pd.DataFrame(
        [
            {"항목": "source_batch_files", "값": len(BATCH_STEMS)},
            {"항목": "raw_rows_before_dedup", "값": raw_before},
            {"항목": "raw_rows_after_url_dedup", "값": len(labeled)},
            {"항목": "모집글적합", "값": int(labeled["모집글적합"].sum())},
            {"항목": "작성자후보적합", "값": int(labeled["작성자후보적합"].sum())},
            {"항목": "학습권장_rows", "값": len(training)},
            {"항목": "top_rows", "값": len(top_posts)},
            {"항목": "low_rows", "값": len(low_posts)},
            {"항목": "strict_candidate_handles", "값": len(strict_handles)},
            {"항목": "상세글 진입 여부", "값": "No. 검색 카드 DOM만 수집"},
            {"항목": "프로필 경로 판단", "값": "비로그인 프로필은 계정당 약 4글로 얕아서 이번 통합본에서는 제외"},
        ]
    )
    notes = pd.DataFrame(
        [
            {
                "항목": "이번 방향",
                "값": "모임 작성자를 프로필에서 깊게 긁는 대신, 검색 카드에서 모집/모임 의도가 강한 글을 많이 모으고 성과 top/low를 라벨링했습니다.",
            },
            {
                "항목": "ML 우선 사용",
                "값": "ml_training_candidates 시트를 먼저 쓰고, 중간 성과까지 회귀/랭킹에 넣으려면 discovery_posts_raw를 함께 쓰면 됩니다.",
            },
            {
                "항목": "라벨 기준",
                "값": "모집글적합=True 중 global 상위 20% 또는 키워드별 상위 20%는 top, global 하위 40% 또는 키워드별 하위 40%는 low로 표시했습니다.",
            },
            {
                "항목": "주의",
                "값": "검색어가 정확히 일치해야 해서 죽은 seed가 있습니다. seed_summary에서 0건 키워드를 제외하고 다음 라운드 seed를 고르면 됩니다.",
            },
        ]
    )

    output_path = run_dir / f"{OUTPUT_STEM}.xlsx"
    jsonl_path = run_dir / f"{OUTPUT_STEM}.jsonl"
    with jsonl_path.open("w", encoding="utf-8") as handle:
        for row in candidates.to_dict(orient="records"):
            handle.write(json.dumps(row, ensure_ascii=False, default=str) + "\n")

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        notes.to_excel(writer, sheet_name="research_notes", index=False)
        quality.to_excel(writer, sheet_name="quality_summary", index=False)
        training.to_excel(writer, sheet_name="ml_training_candidates", index=False)
        top_posts.to_excel(writer, sheet_name="top_posts", index=False)
        low_posts.to_excel(writer, sheet_name="low_posts", index=False)
        strict_handles.to_excel(writer, sheet_name="strict_candidate_handles", index=False)
        seed_summary.to_excel(writer, sheet_name="seed_summary", index=False)
        labeled.to_excel(writer, sheet_name="discovery_posts_raw", index=False)
        pd.DataFrame(sources).to_excel(writer, sheet_name="source_files", index=False)

    format_workbook(output_path)

    print(f"output_excel={output_path}")
    print(f"output_jsonl={jsonl_path}")
    print("\nquality_summary")
    print(quality.to_string(index=False))
    print("\nseed_summary")
    print(seed_summary[["검색어", "수집글수", "후보글수", "작성자수"]].to_string(index=False))
    print("\nstrict_candidate_handles_top")
    print(strict_handles.head(20).to_string(index=False))


if __name__ == "__main__":
    main()
