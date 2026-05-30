from pathlib import Path
import json
import sys

import pandas as pd
from openpyxl import load_workbook


try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


RUNS = [
    ("batch1", "threads_keyword_research_20260514_155456", "친구없음, 외롭다, 인간관계, 소모임, 소개팅앱"),
    ("batch2", "threads_keyword_research_20260514_155811", "내향인, 오픈채팅, 스몰토크, 동호회, 공허함"),
    ("batch3", "threads_keyword_research_20260514_160159", "외로운, 외로워, 혼자, 혼밥, 1인가구"),
    ("batch4", "threads_keyword_research_20260514_160603", "주말, 관계, 모임, 소개팅, 외로움"),
    ("retry_friend", "threads_keyword_research_20260514_161038", "친구없음 단독 재시도"),
]


LABEL_SCHEMA = [
    ("라벨상태", "blank / reviewed / exclude", "라벨링 상태"),
    ("전략보존", "Y / N", "전략 추출 후보 여부"),
    ("적합도점수", "0-5", "사업 핵심 문제 적합도"),
    ("핵심문제라벨", "free text", "강한 문제 프레임"),
    ("세그먼트추정", "free text", "직장인, 내향인, 1인가구 등"),
    ("마케팅각도", "free text", "공감형, 불만형, 오퍼형 등"),
    ("외로움직접표현", "0/1", "외로움 직접 표현"),
    ("사람만나기시도", "0/1", "모임/소개팅/오픈채팅 등 시도"),
    ("모임실패경험", "0/1", "모임 불만/실패 경험"),
    ("관계지속실패", "0/1", "후속 연락/관계 지속 실패"),
    ("스몰토크불만", "0/1", "얕은 대화 불만"),
    ("깊은대화욕구", "0/1", "진짜 이야기/깊은 대화 욕구"),
    ("기존대안피로", "0/1", "기존 서비스/대안 피로"),
    ("첫참석불안", "0/1", "첫 참석/어색함 불안"),
    ("안전품질불안", "0/1", "이상한 사람/품질/안전 우려"),
    ("후속연락욕구", "0/1", "다시 연락할 사람/명분 욕구"),
    ("행동의향", "0/1", "참여/신청/궁금 등 행동 신호"),
    ("경쟁서비스언급", "free text", "문토, 타임레프트 등"),
    ("라벨근거메모", "free text", "라벨 판단 근거"),
]


def summarize_text(value, limit=180):
    text = " ".join(str(value or "").split())
    if len(text) <= limit:
        return text
    return text[: limit - 1] + "…"


def main():
    run_dir = Path(__file__).resolve().parent / "keyword_research_runs"
    output_path = run_dir / "threads_keyword_research_deep_20260514.xlsx"
    jsonl_path = output_path.with_suffix(".jsonl")

    raw_frames = []
    summary_frames = []
    batch_rows = []
    for batch_name, stem, note in RUNS:
        path = run_dir / f"{stem}.xlsx"
        sheets = pd.read_excel(path, sheet_name=None)
        raw = sheets.get("raw_posts", pd.DataFrame()).copy()
        summary = sheets.get("keyword_summary", pd.DataFrame()).copy()
        raw["배치"] = batch_name
        raw["소스파일"] = path.name
        summary["배치"] = batch_name
        summary["배치메모"] = note
        raw_frames.append(raw)
        summary_frames.append(summary)
        batch_rows.append(
            {
                "배치": batch_name,
                "파일": path.name,
                "키워드": note,
                "raw_posts": len(raw),
                "키워드수": len(summary),
            }
        )

    raw_all = pd.concat(raw_frames, ignore_index=True)
    summary_all = pd.concat(summary_frames, ignore_index=True)
    raw_all["본문길이"] = raw_all["본문_카드표시"].fillna("").astype(str).str.len()
    raw_all["통합중복URL여부"] = raw_all["게시글링크"].duplicated(keep=False)
    before = len(raw_all)
    raw = raw_all.drop_duplicates(subset=["게시글링크"], keep="last").copy()
    after_url = len(raw)
    raw = raw[raw["본문_카드표시"].fillna("").astype(str).str.strip() != ""].copy()
    raw["전체반응순위"] = raw["반응합계"].rank(method="first", ascending=False).astype(int)
    raw["키워드내반응순위"] = raw.groupby("키워드")["반응합계"].rank(method="first", ascending=False).astype(int)

    keyword_quality = (
        raw.groupby(["키워드", "검색쿼리"], dropna=False)
        .agg(
            유효게시글수=("게시글링크", "count"),
            평균반응합계=("반응합계", "mean"),
            최대반응합계=("반응합계", "max"),
            댓글합계=("댓글수", "sum"),
            평균본문길이=("본문길이", "mean"),
        )
        .reset_index()
    )
    keyword_quality["평균반응합계"] = keyword_quality["평균반응합계"].round(1)
    keyword_quality["평균본문길이"] = keyword_quality["평균본문길이"].round(1)
    keyword_quality = keyword_quality.sort_values(["유효게시글수", "평균반응합계"], ascending=[False, False])

    top_posts = raw.sort_values(["반응합계", "댓글수"], ascending=[False, False]).head(60).copy()
    top_posts["본문요약"] = top_posts["본문_카드표시"].map(summarize_text)

    quality_summary = pd.DataFrame(
        [
            {"항목": "통합 Excel", "값": str(output_path)},
            {"항목": "통합 JSONL", "값": str(jsonl_path)},
            {"항목": "입력 배치 수", "값": len(RUNS)},
            {"항목": "검색 키워드 실행 수", "값": len(summary_all)},
            {"항목": "중복 제거 전 raw post 수", "값": before},
            {"항목": "URL 중복 제거 후 raw post 수", "값": after_url},
            {"항목": "유효 본문 raw post 수", "값": len(raw)},
            {"항목": "중복 URL 제거 수", "값": before - after_url},
            {"항목": "빈 본문 제외 수", "값": after_url - len(raw)},
            {"항목": "게시글이 1개 이상 나온 키워드 수", "값": raw["키워드"].nunique()},
            {"항목": "rate_limited 키워드 수", "값": int((summary_all["중단이유"] == "rate_limited").sum())},
            {"항목": "no_results 키워드 수", "값": int((summary_all["중단이유"] == "no_results").sum())},
            {"항목": "스크롤 수정", "값": "window가 아니라 Threads 내부 overflow-y:auto 컨테이너를 스크롤"},
            {"항목": "상세글 진입 여부", "값": "No. 검색 결과 카드 DOM만 수집"},
        ]
    )

    with jsonl_path.open("w", encoding="utf-8") as handle:
        for row in raw.to_dict(orient="records"):
            handle.write(json.dumps(row, ensure_ascii=False, default=str) + "\n")

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        quality_summary.to_excel(writer, sheet_name="quality_summary", index=False)
        pd.DataFrame(batch_rows).to_excel(writer, sheet_name="batch_summary", index=False)
        keyword_quality.to_excel(writer, sheet_name="keyword_quality", index=False)
        top_posts.to_excel(writer, sheet_name="top_posts", index=False)
        raw.to_excel(writer, sheet_name="raw_posts", index=False)
        summary_all.to_excel(writer, sheet_name="keyword_summary", index=False)
        pd.DataFrame(LABEL_SCHEMA, columns=["label", "value", "description"]).to_excel(
            writer, sheet_name="label_schema", index=False
        )

    format_workbook(output_path)

    print("quality_summary")
    print(quality_summary.to_string(index=False))
    print("\nkeyword_quality")
    print(keyword_quality.to_string(index=False))


def format_workbook(workbook_path):
    wb = load_workbook(workbook_path)
    widths = {
        "quality_summary": {"A": 28, "B": 120},
        "batch_summary": {"A": 16, "B": 48, "C": 70},
        "keyword_quality": {"A": 18, "B": 18},
        "top_posts": {"J": 80, "K": 70, "L": 80},
        "raw_posts": {"I": 70, "N": 70, "P": 50, "AN": 24, "AO": 20, "AP": 20},
        "label_schema": {"A": 22, "B": 26, "C": 70},
    }
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row and ws.max_column:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.style = "Headline 4"
        for col, width in widths.get(ws.title, {}).items():
            ws.column_dimensions[col].width = width
    wb.save(workbook_path)


if __name__ == "__main__":
    main()
