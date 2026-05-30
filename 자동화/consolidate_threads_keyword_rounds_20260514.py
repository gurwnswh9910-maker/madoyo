from pathlib import Path
import json
import sys

import pandas as pd
from openpyxl import load_workbook


try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


RUNS = [
    {
        "round": 1,
        "stem": "threads_keyword_research_20260514_152200",
        "decision": "외롭- 감정 표현과 친구 없음 변형을 먼저 검증",
        "next_basis": "외롭다/외로운/외로워는 작동, 친구없/친구 없는은 실패",
    },
    {
        "round": 2,
        "stem": "threads_keyword_research_20260514_152336",
        "decision": "친구 없음 표현을 실제 표기 변형으로 재검증",
        "next_basis": "친구없음만 작동, 조사 포함 문장형은 실패",
    },
    {
        "round": 3,
        "stem": "threads_keyword_research_20260514_152451",
        "decision": "짧고 넓은 관계/모임 명사형으로 수집량 확보",
        "next_basis": "명사형 토큰은 전반적으로 안정적",
    },
    {
        "round": 4,
        "stem": "threads_keyword_research_20260514_152721",
        "decision": "외롭- 계열 긴 활용형을 추가 검증",
        "next_basis": "긴 활용형은 거의 실패, 짧은 토큰으로 회귀 필요",
    },
    {
        "round": 5,
        "stem": "threads_keyword_research_20260514_152914",
        "decision": "생활/고립 맥락의 짧은 토큰으로 수집량 확보",
        "next_basis": "혼자/혼밥/주말/공허함/1인가구 모두 안정적으로 4건씩 수집",
    },
    {
        "round": 6,
        "stem": "threads_keyword_research_20260514_153133",
        "decision": "기존 대안과 대화 불만 토큰으로 마무리 확장",
        "next_basis": "소개팅앱/오픈채팅/스몰토크/동호회가 안정적",
    },
]


LABEL_SCHEMA = [
    ("라벨상태", "blank / reviewed / exclude", "라벨링 상태"),
    ("전략보존", "Y / N", "전략 추출 후보 여부"),
    ("적합도점수", "0-5", "사업 핵심 문제 적합도"),
    ("핵심문제라벨", "free text", "강한 문제 프레임"),
    ("세그먼트추정", "free text", "직장인, 내향인, 1인가구 등"),
    ("마케팅각도", "free text", "공감형, 불만형, 오퍼형 등"),
    ("외로움직접표현", "0/1", "외로움 직접 표현"),
    ("사람만나기시도", "0/1", "모임/소개팅/오픈채팅 등 시도"),
    ("모임실패경험", "0/1", "모임 불만/실패 경험"),
    ("관계지속실패", "0/1", "후속 연락/관계 지속 실패"),
    ("스몰토크불만", "0/1", "얕은 대화 불만"),
    ("깊은대화욕구", "0/1", "진짜 이야기/깊은 대화 욕구"),
    ("기존대안피로", "0/1", "기존 서비스/대안 피로"),
    ("첫참석불안", "0/1", "첫 참석/어색함 불안"),
    ("안전품질불안", "0/1", "이상한 사람/품질/안전 우려"),
    ("후속연락욕구", "0/1", "다시 연락할 사람/명분 욕구"),
    ("행동의향", "0/1", "참여/신청/궁금 등 행동 신호"),
    ("경쟁서비스언급", "free text", "문토, 타임레프트 등"),
    ("라벨근거메모", "free text", "라벨 판단 근거"),
]


def summarize_text(value, limit=160):
    text = " ".join(str(value or "").split())
    if len(text) <= limit:
        return text
    return text[: limit - 1] + "…"


def main():
    run_dir = Path(__file__).resolve().parent / "keyword_research_runs"
    output_path = run_dir / "threads_keyword_research_iterative_20260514_6rounds.xlsx"
    jsonl_path = output_path.with_suffix(".jsonl")

    raw_frames = []
    summary_frames = []
    for run in RUNS:
        path = run_dir / f"{run['stem']}.xlsx"
        sheets = pd.read_excel(path, sheet_name=None)
        raw = sheets.get("raw_posts", pd.DataFrame()).copy()
        summary = sheets.get("keyword_summary", pd.DataFrame()).copy()
        raw["라운드"] = run["round"]
        raw["소스파일"] = path.name
        summary["라운드"] = run["round"]
        summary["라운드판단"] = run["decision"]
        summary["다음선정근거"] = run["next_basis"]
        raw_frames.append(raw)
        summary_frames.append(summary)

    raw_all = pd.concat(raw_frames, ignore_index=True) if raw_frames else pd.DataFrame()
    summary_all = pd.concat(summary_frames, ignore_index=True) if summary_frames else pd.DataFrame()
    raw_all["본문길이"] = raw_all["본문_카드표시"].fillna("").astype(str).str.len()
    raw_all["통합중복URL여부"] = raw_all["게시글링크"].duplicated(keep=False)
    before_dedupe = len(raw_all)
    raw_deduped = raw_all.drop_duplicates(subset=["게시글링크"], keep="first").copy()
    after_url_dedupe = len(raw_deduped)
    before_body_filter = len(raw_deduped)
    raw_deduped = raw_deduped[raw_deduped["본문_카드표시"].fillna("").astype(str).str.strip() != ""].copy()
    raw_deduped["전체반응순위"] = raw_deduped["반응합계"].rank(method="first", ascending=False).astype(int)
    raw_deduped["키워드내반응순위"] = raw_deduped.groupby("키워드")["반응합계"].rank(method="first", ascending=False).astype(int)

    round_summary = (
        summary_all.groupby("라운드", dropna=False)
        .agg(
            라운드키워드수=("키워드", "count"),
            실행수집게시글수=("수집게시글수", "sum"),
            rate_limited수=("중단이유", lambda s: int((s == "rate_limited").sum())),
            no_results수=("중단이유", lambda s: int((s == "no_results").sum())),
        )
        .reset_index()
    )
    decisions = pd.DataFrame(RUNS)[["round", "decision", "next_basis"]].rename(
        columns={"round": "라운드", "decision": "라운드판단", "next_basis": "다음선정근거"}
    )
    round_summary = round_summary.merge(decisions, on="라운드", how="left")

    keyword_quality = (
        raw_deduped.groupby(["키워드", "검색쿼리"], dropna=False)
        .agg(
            통합게시글수=("게시글링크", "count"),
            평균반응합계=("반응합계", "mean"),
            최대반응합계=("반응합계", "max"),
            평균본문길이=("본문길이", "mean"),
            댓글합계=("댓글수", "sum"),
        )
        .reset_index()
    )
    keyword_quality["평균반응합계"] = keyword_quality["평균반응합계"].round(1)
    keyword_quality["평균본문길이"] = keyword_quality["평균본문길이"].round(1)
    keyword_quality = keyword_quality.sort_values(["통합게시글수", "평균반응합계"], ascending=[False, False])

    top_posts = raw_deduped.sort_values(["반응합계", "댓글수"], ascending=[False, False]).head(40).copy()
    top_posts["본문요약"] = top_posts["본문_카드표시"].map(summarize_text)

    quality_summary = pd.DataFrame(
        [
            {"항목": "통합 Excel", "값": str(output_path)},
            {"항목": "통합 JSONL", "값": str(jsonl_path)},
            {"항목": "라운드 수", "값": len(RUNS)},
            {"항목": "검색 키워드 수", "값": len(summary_all)},
            {"항목": "중복 제거 전 raw post 수", "값": before_dedupe},
            {"항목": "URL 중복 제거 후 raw post 수", "값": after_url_dedupe},
            {"항목": "유효 본문 raw post 수", "값": len(raw_deduped)},
            {"항목": "빈 본문 제외 수", "값": before_body_filter - len(raw_deduped)},
            {"항목": "게시글이 1개 이상 나온 키워드 수", "값": raw_deduped["키워드"].nunique()},
            {"항목": "중복 URL 제거 수", "값": before_dedupe - after_url_dedupe},
            {"항목": "빈 본문 행 수", "값": int((raw_deduped["본문_카드표시"].fillna("").astype(str).str.strip() == "").sum())},
            {"항목": "rate_limited 키워드 수", "값": int((summary_all["중단이유"] == "rate_limited").sum())},
            {"항목": "no_results 키워드 수", "값": int((summary_all["중단이유"] == "no_results").sum())},
            {"항목": "상세글 진입 여부", "값": "No. 검색 결과 카드 DOM만 수집"},
            {"항목": "운영 결론", "값": "Threads 비로그인 검색은 키워드당 공개 카드 수가 작으므로, 짧은 정확 문자열을 넓게 쪼개 총량을 확보해야 함"},
        ]
    )

    next_keywords = pd.DataFrame(
        [
            ("외롭다", "유지", "외롭- 계열 중 가장 안정적으로 수집"),
            ("외로운", "유지", "감정 표현 카드 수집 가능"),
            ("외로워", "유지", "구어체 감정 표현"),
            ("친구없음", "유지", "친구 없음 자기 고백형이 강함"),
            ("인간관계", "유지", "관계 고민 전반"),
            ("내향인", "유지", "댓글 반응과 세그먼트성이 좋음"),
            ("소모임", "유지", "기존 대안/모임 맥락"),
            ("소개팅앱", "유지", "기존 대안 피로"),
            ("오픈채팅", "유지", "안전/품질 불안 후보"),
            ("스몰토크", "유지", "대화 불만 축"),
            ("외로운 사람", "2차 실험", "문장형이지만 오퍼 문구와 가까움"),
            ("모임 어색", "2차 실험", "공백 포함이 약할 수 있으므로 별도 검증"),
            ("소모임 어색", "2차 실험", "문제 가설과 직접 연결"),
            ("친구 사귀기", "2차 실험", "행동 의향 검색어"),
            ("사람 만나는", "2차 실험", "표현형 검증 필요"),
        ],
        columns=["키워드", "권장", "이유"],
    )

    with jsonl_path.open("w", encoding="utf-8") as handle:
        for row in raw_deduped.to_dict(orient="records"):
            handle.write(json.dumps(row, ensure_ascii=False, default=str) + "\n")

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        quality_summary.to_excel(writer, sheet_name="quality_summary", index=False)
        round_summary.to_excel(writer, sheet_name="round_summary", index=False)
        keyword_quality.to_excel(writer, sheet_name="keyword_quality", index=False)
        top_posts.to_excel(writer, sheet_name="top_posts", index=False)
        next_keywords.to_excel(writer, sheet_name="next_keywords", index=False)
        raw_deduped.to_excel(writer, sheet_name="raw_posts", index=False)
        summary_all.to_excel(writer, sheet_name="keyword_summary", index=False)
        pd.DataFrame(LABEL_SCHEMA, columns=["label", "value", "description"]).to_excel(
            writer, sheet_name="label_schema", index=False
        )

    format_workbook(output_path)

    print("quality_summary")
    print(quality_summary.to_string(index=False))
    print("\nround_summary")
    print(round_summary.to_string(index=False))
    print("\nkeyword_quality_top")
    print(keyword_quality.head(15).to_string(index=False))
    print("\ntop_posts")
    print(top_posts[["전체반응순위", "라운드", "키워드", "작성자", "반응합계", "댓글수", "본문요약"]].head(15).to_string(index=False))


def format_workbook(workbook_path):
    wb = load_workbook(workbook_path)
    widths = {
        "quality_summary": {"A": 26, "B": 120},
        "round_summary": {"A": 10, "F": 60, "G": 70},
        "keyword_quality": {"A": 18, "B": 18},
        "top_posts": {"J": 80, "K": 70, "L": 80},
        "next_keywords": {"A": 18, "B": 16, "C": 70},
        "raw_posts": {"I": 70, "N": 70, "P": 50, "AN": 24, "AO": 20, "AP": 20},
        "label_schema": {"A": 22, "B": 26, "C": 70},
    }
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row and ws.max_column:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.style = "Headline 4"
        for col, width in widths.get(ws.title, {}).items():
            ws.column_dimensions[col].width = width
    wb.save(workbook_path)


if __name__ == "__main__":
    main()
