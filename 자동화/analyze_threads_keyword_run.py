from pathlib import Path
import sys

import pandas as pd
from openpyxl import load_workbook


RUN_STEM = "threads_keyword_research_20260514_145546"

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


def summarize_text(value, limit=160):
    text = " ".join(str(value or "").split())
    if len(text) <= limit:
        return text
    return text[: limit - 1] + "…"


def is_body_noise(text, handle=""):
    value = " ".join(str(text or "").split()).strip()
    if not value:
        return True
    handle_value = str(handle or "").strip().lstrip("@")
    if handle_value and value in {handle_value, "@" + handle_value}:
        return True
    if value in {"번역하기", "See translation", "좋아요", "답글", "댓글", "리포스트", "공유", "공유하기"}:
        return True
    if value.replace(".", "", 1).replace(",", "").isdigit():
        return True
    if value.endswith("천") or value.endswith("만"):
        return any(ch.isdigit() for ch in value)
    if len(value) <= 5 and any(ch.isdigit() for ch in value):
        return True
    if len(value) == 10 and value[4] == "-" and value[7] == "-":
        return True
    return False


def repair_body_from_blocks(row):
    current = str(row.get("본문_카드표시", "") or "").strip()
    handle = str(row.get("작성자", "") or "")
    if current and not is_body_noise(current, handle):
        return current
    try:
        blocks = json_loads(row.get("본문_후보블록", "[]"))
    except Exception:
        blocks = []
    candidates = [str(block).strip() for block in blocks if not is_body_noise(block, handle)]
    if candidates:
        return sorted(candidates, key=len, reverse=True)[0]
    return "" if is_body_noise(current, handle) else current


def json_loads(value):
    import json

    if pd.isna(value):
        return []
    return json.loads(str(value))


def main():
    run_dir = Path(__file__).resolve().parent / "keyword_research_runs"
    workbook_path = run_dir / f"{RUN_STEM}.xlsx"
    jsonl_path = run_dir / f"{RUN_STEM}.jsonl"
    sheets = pd.read_excel(workbook_path, sheet_name=None)
    raw = sheets["raw_posts"].copy()
    summary = sheets["keyword_summary"].copy()

    raw["본문_카드표시"] = raw.apply(repair_body_from_blocks, axis=1)
    raw["본문길이"] = raw["본문_카드표시"].fillna("").astype(str).str.len()
    raw["중복URL여부"] = raw["게시글링크"].duplicated(keep=False)
    raw["전체반응순위"] = raw["반응합계"].rank(method="first", ascending=False).astype(int)
    raw["키워드내반응순위"] = raw.groupby("키워드")["반응합계"].rank(method="first", ascending=False).astype(int)

    keyword_quality = (
        raw.groupby(["키워드", "검색쿼리", "키워드분류"], dropna=False)
        .agg(
            수집게시글수=("게시글링크", "count"),
            평균반응합계=("반응합계", "mean"),
            최대반응합계=("반응합계", "max"),
            평균본문길이=("본문길이", "mean"),
            외부링크글수=("외부링크있음", "sum"),
        )
        .reset_index()
    )
    keyword_quality["평균반응합계"] = keyword_quality["평균반응합계"].round(1)
    keyword_quality["평균본문길이"] = keyword_quality["평균본문길이"].round(1)
    keyword_quality = summary.merge(
        keyword_quality,
        on=["키워드", "검색쿼리", "키워드분류"],
        how="left",
    )
    keyword_quality["수집게시글수_y"] = keyword_quality["수집게시글수_y"].fillna(0).astype(int)
    keyword_quality = keyword_quality.rename(columns={"수집게시글수_x": "실행수집게시글수", "수집게시글수_y": "raw_posts행수"})

    top_cols = [
        "전체반응순위",
        "키워드",
        "검색쿼리",
        "작성자",
        "반응합계",
        "좋아요수",
        "댓글수",
        "리포스트수",
        "공유수",
        "본문_카드표시",
        "게시글링크",
    ]
    top_posts = raw.sort_values(["반응합계", "댓글수"], ascending=[False, False])[top_cols].head(30).copy()
    top_posts["본문요약"] = top_posts["본문_카드표시"].map(summarize_text)

    quality_summary = pd.DataFrame(
        [
            {"항목": "Excel 결과", "값": str(workbook_path)},
            {"항목": "JSONL 원본", "값": str(jsonl_path)},
            {"항목": "총 키워드 수", "값": len(summary)},
            {"항목": "수집된 raw post 수", "값": len(raw)},
            {"항목": "게시글이 1개 이상 나온 키워드 수", "값": raw["키워드"].nunique()},
            {"항목": "중복 URL 행 수", "값": int(raw["게시글링크"].duplicated().sum())},
            {"항목": "빈 본문 행 수", "값": int((raw["본문_카드표시"].fillna("").astype(str).str.strip() == "").sum())},
            {"항목": "rate_limited 키워드 수", "값": int((summary["중단이유"] == "rate_limited").sum())},
            {"항목": "no_results 키워드 수", "값": int((summary["중단이유"] == "no_results").sum())},
            {"항목": "idle_limit 키워드 수", "값": int((summary["중단이유"] == "idle_limit_reached").sum())},
            {"항목": "상세글 진입 여부", "값": "No. 검색 결과 카드 DOM만 수집"},
            {"항목": "검증 메모", "값": "Threads /search 입력창 Enter 결과가 serp_type=default이며, 이 경로에서 카드 DOM 수집 가능"},
        ]
    )

    ordered = {
        "quality_summary": quality_summary,
        "keyword_quality": keyword_quality,
        "top_posts": top_posts,
        "raw_posts": raw,
        "keyword_summary": summary,
        "label_schema": sheets["label_schema"],
        "run_config": sheets["run_config"],
    }

    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        for sheet_name, df in ordered.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    format_workbook(workbook_path)

    print("quality_summary")
    print(quality_summary.to_string(index=False))
    print("\nkeyword_quality_top")
    print(keyword_quality.sort_values("raw_posts행수", ascending=False).head(12).to_string(index=False))
    print("\ntop_posts")
    print(top_posts[["전체반응순위", "키워드", "작성자", "반응합계", "댓글수", "본문요약"]].head(12).to_string(index=False))


def format_workbook(workbook_path):
    wb = load_workbook(workbook_path)
    widths = {
        "quality_summary": {"A": 24, "B": 120},
        "keyword_quality": {"A": 14, "B": 14, "C": 18, "I": 24},
        "top_posts": {"J": 80, "K": 70, "L": 80},
        "raw_posts": {"I": 70, "N": 70, "P": 50, "AN": 24, "AO": 20, "AP": 20},
        "label_schema": {"A": 22, "B": 26, "C": 70},
    }
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row and ws.max_column:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.style = "Headline 4"
        for col, width in widths.get(ws.title, {}).items():
            ws.column_dimensions[col].width = width
    wb.save(workbook_path)


if __name__ == "__main__":
    main()
